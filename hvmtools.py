#!/usr/bin/env python3
# SPDX-License-Identifier: MIT
# Copyright (c) 2026 Graham Fairbank. See LICENSE for details.
"""
HVMTools - an RVTools-style inventory exporter for HPE Morpheus VM Essentials (HVM).

Connects to the VM Essentials manager (Morpheus API), gathers clouds, clusters,
hosts, VMs, instances, disks, NICs, networks, datastores and snapshots, and writes
an Excel workbook with RVTools-like tabs (vInfo, vCPU, vMemory, vDisk, vNetwork,
vSnapshot, vHost, vCluster, vDatastore, vNetworks, vInstance, vHealth, vSource).

Usage:
    pip install requests openpyxl
    python hvmtools.py --url https://hvm-manager.example.com -u admin --insecure
    python hvmtools.py --url https://hvm-manager --token <API_TOKEN> -o inventory.xlsx
    python hvmtools.py --from-json inventory.json            # rebuild offline

Environment variables (optional): MORPHEUS_URL, MORPHEUS_USER, MORPHEUS_PASS, MORPHEUS_TOKEN
"""
import argparse
import datetime as dt
import getpass
import json
import os
import re
import sys
from concurrent.futures import ThreadPoolExecutor

import requests
import urllib3
from openpyxl import Workbook
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

__version__ = "1.0.0"
MIB = 1024 * 1024


# --------------------------------------------------------------------------- helpers
def g(d, *path, default=None):
    """Safe nested get: g(obj, 'a', 'b', 0)."""
    cur = d
    for p in path:
        if isinstance(cur, dict):
            cur = cur.get(p)
        elif isinstance(cur, list) and isinstance(p, int):
            cur = cur[p] if -len(cur) <= p < len(cur) else None
        else:
            return default
        if cur is None:
            return default
    return cur


def first(*vals):
    for v in vals:
        if v not in (None, "", [], {}):
            return v
    return None


def to_int(v):
    try:
        return int(float(v)) if v not in (None, "") else None
    except (TypeError, ValueError):
        return None


def rnd(v, n=1):
    try:
        return round(float(v), n) if v not in (None, "") else None
    except (TypeError, ValueError):
        return None


def mib(b):
    n = to_int(b)
    return round(n / MIB) if n is not None else None


def pct(used, total):
    u, t = to_int(used), to_int(total)
    return round(u / t * 100, 1) if u is not None and t else None


def parse_date(s):
    if not s or not isinstance(s, str):
        return None
    try:
        d = dt.datetime.fromisoformat(s.replace("Z", "+00:00"))
        return d.astimezone(dt.timezone.utc).replace(tzinfo=None)
    except ValueError:
        return s


def power(state):
    return {"on": "poweredOn", "off": "poweredOff", "suspended": "suspended",
            "paused": "suspended"}.get(str(state).lower(), state) if state else None


def fmt_tags(tags):
    if not tags:
        return None
    if isinstance(tags, list):
        out = []
        for t in tags:
            if isinstance(t, dict):
                k, v = t.get("name") or t.get("key"), t.get("value")
                out.append(f"{k}={v}" if v not in (None, "") else str(k))
            else:
                out.append(str(t))
        return "; ".join(out)
    return str(tags)


def strip(v):
    return v.strip() or None if isinstance(v, str) else v


def name_of(obj):
    if isinstance(obj, dict):
        return first(obj.get("name"), obj.get("displayName"), obj.get("username"))
    return obj


def log(msg):
    print(f"[{dt.datetime.now():%H:%M:%S}] {msg}", file=sys.stderr)


SECRET_KEY_RE = re.compile(r"password|passwd|secret|token|apikey|api_key|privatekey|serviceaccess|servicecert",
                           re.IGNORECASE)


def redact(obj):
    """Blank credential-like fields (sshPassword, apiKey, servicePasswordHash, ...) in place."""
    if isinstance(obj, dict):
        for k, v in obj.items():
            if SECRET_KEY_RE.search(str(k)) and v not in (None, "", [], {}):
                obj[k] = "<redacted>"
            else:
                redact(v)
    elif isinstance(obj, list):
        for v in obj:
            redact(v)
    return obj


# --------------------------------------------------------------------------- API client
class MorpheusClient:
    def __init__(self, url, verify=True, timeout=60, page_size=100):
        self.url = url.rstrip("/")
        if not self.url.startswith("http"):
            self.url = "https://" + self.url
        self.timeout = timeout
        self.page_size = page_size
        self.s = requests.Session()
        self.s.verify = verify
        self.s.headers["Accept"] = "application/json"
        if not verify:
            urllib3.disable_warnings(urllib3.exceptions.InsecureRequestWarning)

    def login(self, username, password, client_id="morph-api"):
        r = self.s.post(f"{self.url}/oauth/token",
                        params={"grant_type": "password", "scope": "write", "client_id": client_id},
                        data={"username": username, "password": password},
                        timeout=self.timeout)
        if r.status_code != 200:
            raise SystemExit(f"Login failed ({r.status_code}): {r.text[:300]}")
        try:
            self.set_token(r.json()["access_token"])
        except (ValueError, KeyError):
            raise SystemExit(f"Login response from {self.url} had no access token - is this a VM Essentials manager?")

    def set_token(self, token):
        self.s.headers["Authorization"] = f"Bearer {token}"

    def get(self, path, params=None, quiet=False):
        try:
            r = self.s.get(f"{self.url}{path}", params=params, timeout=self.timeout)
        except requests.RequestException as e:
            if quiet:
                log(f"  warn: {path}: {e}")
                return None
            raise
        if r.status_code in (401,):
            raise SystemExit("Unauthorized - check credentials / token.")
        if r.status_code >= 400:
            if quiet:
                log(f"  warn: {path} -> HTTP {r.status_code}")
                return None
            r.raise_for_status()
        try:
            return r.json()
        except ValueError:
            return None

    def list(self, path, key, params=None, quiet=True):
        """Paginate a Morpheus list endpoint (max/offset + meta.total)."""
        out, offset = [], 0
        while True:
            p = dict(params or {}, max=self.page_size, offset=offset)
            data = self.get(path, p, quiet=quiet)
            if not data:
                break
            page = data.get(key) or []
            out.extend(page)
            total = to_int(g(data, "meta", "total"))
            offset += len(page)
            if not page or (total is not None and offset >= total) or len(page) < self.page_size:
                break
        return out


# --------------------------------------------------------------------------- collection
def collect(c, snapshots=True, stats=True, workers=8):
    inv = {"collected": dt.datetime.now(dt.timezone.utc).strftime("%Y-%m-%dT%H:%M:%SZ"), "manager": c.url}
    log("Connecting / ping")
    try:  # not quiet: TLS / connection failures must stop the run, not yield an empty workbook
        inv["ping"] = c.get("/api/ping") or {}
    except requests.exceptions.HTTPError:
        inv["ping"] = {}
    inv["whoami"] = c.get("/api/whoami", quiet=True) or {}

    log("Clouds")
    inv["clouds"] = c.list("/api/zones", "zones") or c.list("/api/clouds", "clouds")
    log("Groups")
    inv["groups"] = c.list("/api/groups", "groups")

    log("Clusters")
    clusters = c.list("/api/clusters", "clusters")
    for cl in clusters:  # list view may omit member servers
        if not cl.get("servers"):
            det = c.get(f"/api/clusters/{cl['id']}", quiet=True) or {}
            cl["servers"] = g(det, "cluster", "servers", default=[])
    inv["clusters"] = clusters

    log("Servers (hosts + VMs)")
    inv["servers"] = c.list("/api/servers", "servers")
    log(f"  {len(inv['servers'])} servers")
    if not inv["servers"]:
        log("  warn: no servers returned - check the user's role has read access to Infrastructure")

    log("Instances")
    inv["instances"] = c.list("/api/instances", "instances")
    log(f"  {len(inv['instances'])} instances")

    log("Networks")
    inv["networks"] = c.list("/api/networks", "networks")

    log("Datastores (cluster Storage tabs + clouds)")
    ds_by_id = {}

    def merge_ds(d, **extra):
        cur = ds_by_id.setdefault(d.get("id"), {})
        for k, v in d.items():
            if cur.get(k) in (None, "", [], {}):
                cur[k] = v
        for k, v in extra.items():
            if v is not None:
                cur.setdefault(k, v)
        cur.setdefault("_clusters", [])
        return cur

    # Cluster > Storage > Data Stores (GFS2, OCFS2, NFS pools, local, etc.)
    for cl in inv["clusters"]:
        for d in c.list(f"/api/clusters/{cl['id']}/datastores", "datastores"):
            cur = merge_ds(d, zone=first(cl.get("zone"), cl.get("cloud")))
            if cl.get("name") not in cur["_clusters"]:
                cur["_clusters"].append(cl.get("name"))
    # Cloud-level datastores (catches anything not attached to a cluster)
    for cl in inv["clouds"]:
        for d in c.list(f"/api/zones/{cl['id']}/data-stores", "datastores"):
            merge_ds(d, zone={"id": cl["id"], "name": cl.get("name")})
    inv["datastores"] = list(ds_by_id.values())
    log(f"  {len(inv['datastores'])} datastores")

    # The /api/servers list view can omit usage stats; the detail view has them.
    need = [s for s in inv["servers"]
            if first(g(s, "stats", "usedMemory"), s.get("usedMemory")) is None]
    if stats and need:
        log(f"Usage stats ({len(need)} servers without stats in list view)")

        def detail(s):
            d = c.get(f"/api/servers/{s['id']}", quiet=True) or {}
            merged = {}
            for src in (g(d, "server", "stats"), d.get("stats")):
                for k, v in (src or {}).items():
                    if merged.get(k) in (None, ""):
                        merged[k] = v
            return s, merged

        with ThreadPoolExecutor(max_workers=workers) as ex:
            got = 0
            for s, merged in ex.map(detail, need):
                cur = s.get("stats") if isinstance(s.get("stats"), dict) else {}
                for k, v in merged.items():
                    if cur.get(k) in (None, ""):
                        cur[k] = v
                if cur:
                    s["stats"] = cur
                got += 1 if cur.get("usedMemory") is not None else 0
        log(f"  {got} of {len(need)} returned memory usage")

    inv["snapshots"] = {}
    if snapshots and inv["instances"]:
        log(f"Snapshots ({len(inv['instances'])} instances)")

        def snap(inst):
            d = c.get(f"/api/instances/{inst['id']}/snapshots", quiet=True) or {}
            return str(inst["id"]), d.get("snapshots") or []

        with ThreadPoolExecutor(max_workers=workers) as ex:
            for iid, snaps in ex.map(snap, inv["instances"]):
                if snaps:
                    inv["snapshots"][iid] = snaps
    return redact(inv)


# --------------------------------------------------------------------------- model building
def build_sheets(inv, snapshot_age_days=7, ds_free_pct=10):
    servers = inv.get("servers", [])
    instances = inv.get("instances", [])
    clusters = inv.get("clusters", [])
    manager = inv.get("manager")

    srv_by_id = {s.get("id"): s for s in servers}
    ds_by_id = {d.get("id"): d for d in inv.get("datastores", [])}
    net_by_id = {n.get("id"): n for n in inv.get("networks", [])}
    group_by_id = {gr.get("id"): gr for gr in inv.get("groups", [])}

    # host classification
    parent_ids = {g(s, "parentServer", "id") for s in servers} - {None}
    host_ids = {s.get("id") for s in servers
                if g(s, "computeServerType", "vmHypervisor") is True
                or g(s, "computeServerType", "format") == "vmHypervisor"} | parent_ids
    hosts = [s for s in servers if s.get("id") in host_ids]
    vms = [s for s in servers if s.get("id") not in host_ids]

    # cluster membership
    host_cluster = {}
    for cl in clusters:
        for m in cl.get("servers") or []:
            sid = m.get("id") if isinstance(m, dict) else m
            host_cluster[sid] = cl.get("name")

    def cluster_of(s):
        return first(name_of(s.get("cluster")), name_of(s.get("serverGroup")),
                     host_cluster.get(s.get("id")),
                     host_cluster.get(g(s, "parentServer", "id")))

    # server -> instance
    inst_by_server = {}
    for i in instances:
        for sid in i.get("servers") or []:
            inst_by_server[sid.get("id") if isinstance(sid, dict) else sid] = i
        for ct in i.get("containers") or []:
            sid = g(ct, "server", "id") if isinstance(ct, dict) else None
            if sid:
                inst_by_server.setdefault(sid, i)

    inst_servers = {}
    for sid, i in inst_by_server.items():
        inst_servers[i.get("id")] = inst_servers.get(i.get("id"), 0) + 1

    def stat(s, inst, *keys):
        """Usage stat from the server, else from its instance (single-VM instances only)."""
        vals = [g(s, "stats", k) for k in keys] + [s.get(keys[0])]
        if inst and inst_servers.get(inst.get("id")) == 1:
            vals += [g(inst, "stats", k) for k in keys]
        return first(*vals)

    def host_name(s):
        return first(g(s, "parentServer", "name"), name_of(srv_by_id.get(g(s, "parentServer", "id"))))

    def cloud_name(s):
        return first(g(s, "zone", "name"), g(s, "cloud", "name"))

    def ds_name(v):
        dsid = first(v.get("datastoreId"), g(v, "datastore", "id"))
        return first(g(v, "datastore", "name"), name_of(ds_by_id.get(dsid)))

    def ips(s):
        vals = [s.get("internalIp"), s.get("externalIp")]
        vals += [n.get("ipAddress") for n in s.get("interfaces") or []]
        seen = []
        for v in vals:
            if v and v not in seen:
                seen.append(v)
        return seen

    rows = {k: [] for k in ["vInfo", "vCPU", "vMemory", "vDisk", "vNetwork", "vSnapshot",
                            "vHost", "vCluster", "vDatastore", "vNetworks", "vInstance",
                            "vHealth", "vSource"]}
    health = rows["vHealth"]

    # ---- VMs
    host_vm_stats = {}
    for s in vms:
        inst = inst_by_server.get(s.get("id"), {})
        vols = s.get("volumes") or []
        nics = s.get("interfaces") or []
        ps = power(s.get("powerState"))
        cores = to_int(first(s.get("maxCores"), inst.get("maxCores"), g(inst, "plan", "maxCores")))
        cps = to_int(first(s.get("coresPerSocket"), inst.get("coresPerSocket"))) or None
        mem = to_int(first(s.get("maxMemory"), inst.get("maxMemory")))
        used_mem = stat(s, inst, "usedMemory")
        cpu_use = stat(s, inst, "cpuUsage", "usedCpu")
        cluster = cluster_of(s)
        host = host_name(s)
        common = {"Host": host, "Cluster": cluster, "Cloud": cloud_name(s), "VM ID": s.get("id")}
        netnames = [first(g(n, "network", "name"), name_of(net_by_id.get(g(n, "network", "id"))))
                    for n in nics]
        ip_list = ips(s)
        group = first(g(inst, "group", "name"), g(s, "site", "name"),
                      name_of(group_by_id.get(g(s, "site", "id"))))

        rows["vInfo"].append({
            "VM": first(s.get("displayName"), s.get("name")), "Powerstate": ps,
            "Status": s.get("status"), "Template": False,
            "DNS Name": first(s.get("hostname"), s.get("externalFqdn")),
            "Agent Installed": s.get("agentInstalled"), "Agent Version": s.get("agentVersion"),
            "Creation date": parse_date(s.get("dateCreated")),
            "Last Updated": parse_date(s.get("lastUpdated")),
            "CPUs": cores, "Sockets": (cores // cps) if cores and cps else None,
            "Cores per Socket": cps, "Memory MiB": mib(mem), "NICs": len(nics), "Disks": len(vols),
            "Total disk capacity MiB": mib(sum(to_int(v.get("maxStorage")) or 0 for v in vols)) if vols else None,
            "Provisioned MiB": mib(s.get("maxStorage")),
            "In Use MiB": mib(stat(s, inst, "usedStorage")),
            "Primary IP Address": first(s.get("internalIp"), s.get("externalIp"), s.get("sshHost")),
            "All IP Addresses": ", ".join(ip_list) or None,
            "Network #1": g(netnames, 0), "Network #2": g(netnames, 1),
            "Network #3": g(netnames, 2), "Network #4": g(netnames, 3),
            "Resource Pool": g(s, "resourcePool", "name"), "Group": group,
            "Instance": inst.get("name"), "Instance Type": g(inst, "instanceType", "name"),
            "Layout": g(inst, "layout", "name"), "Plan": first(g(s, "plan", "name"), g(inst, "plan", "name")),
            "Environment": inst.get("environment"),
            "Owner": first(g(inst, "owner", "username"), g(inst, "createdBy", "username"),
                           g(s, "owner", "username")),
            "Annotation": first(s.get("description"), inst.get("description")),
            "Cloud": common["Cloud"], "Cluster": cluster, "Host": host,
            "OS": first(g(s, "serverOs", "name"), g(s, "osType")),
            "OS Platform": s.get("platform"), "OS Version": s.get("platformVersion"),
            "Tags": fmt_tags(first(s.get("tags"), inst.get("tags"))),
            "Labels": ", ".join(first(s.get("labels"), inst.get("labels")) or []) or None,
            "Managed": first(s.get("managed"), g(s, "computeServerType", "managed")),
            "Discovered": s.get("discovered"),
            "Server Type": g(s, "computeServerType", "name"),
            "VM ID": s.get("id"), "Instance ID": inst.get("id"),
            "External ID": s.get("externalId"), "UUID": s.get("uuid"), "Manager": manager,
        })

        rows["vCPU"].append({
            "VM": first(s.get("displayName"), s.get("name")), "Powerstate": ps, "CPUs": cores,
            "Sockets": (cores // cps) if cores and cps else None, "Cores per Socket": cps,
            "CPU Usage %": rnd(cpu_use),
            "Plan": g(s, "plan", "name"), "OS": first(g(s, "serverOs", "name"), s.get("osType")), **common})

        rows["vMemory"].append({
            "VM": first(s.get("displayName"), s.get("name")), "Powerstate": ps,
            "Size MiB": mib(mem), "Consumed MiB": mib(used_mem),
            "Free MiB": mib(to_int(mem) - to_int(used_mem)) if mem and used_mem is not None else None,
            "Usage %": pct(used_mem, mem), "Plan": g(s, "plan", "name"), **common})

        for v in vols:
            rows["vDisk"].append({
                "VM": first(s.get("displayName"), s.get("name")), "Powerstate": ps,
                "Disk": v.get("name"), "Capacity MiB": mib(v.get("maxStorage")),
                "Used MiB": mib(v.get("usedStorage")), "Root/Boot": v.get("rootVolume"),
                "Device": first(v.get("deviceDisplayName"), v.get("deviceName")),
                "Controller": first(g(v, "controller", "name"), v.get("controllerId")),
                "Unit #": to_int(v.get("unitNumber")), "Datastore": ds_name(v),
                "Storage Type": first(g(v, "type", "name"), g(v, "storageType", "name"), v.get("typeName")),
                "Thin": v.get("thin"), "Status": v.get("status"), "Disk ID": v.get("id"),
                "External ID": v.get("externalId"), **common})

        for n in nics:
            net = net_by_id.get(g(n, "network", "id"), {})
            rows["vNetwork"].append({
                "VM": first(s.get("displayName"), s.get("name")), "Powerstate": ps,
                "NIC": first(n.get("name"), n.get("displayName")),
                "Network": first(g(n, "network", "name"), net.get("name")),
                "VLAN": net.get("vlanId"), "Primary": n.get("primaryInterface"),
                "IPv4 Address": n.get("ipAddress"), "IPv6 Address": n.get("ipv6Address"),
                "MAC Address": n.get("macAddress"),
                "Adapter": first(g(n, "type", "name"), n.get("type") if isinstance(n.get("type"), str) else None),
                "IP Mode": n.get("ipMode"), "Connected": n.get("active"), "NIC ID": n.get("id"), **common})

        # health
        if ps == "poweredOff":
            health.append({"Entity": s.get("name"), "Type": "VM", "Severity": "Info", "Message": "VM is powered off"})
        if s.get("agentInstalled") is False and ps == "poweredOn":
            health.append({"Entity": s.get("name"), "Type": "VM", "Severity": "Warning",
                           "Message": "Morpheus agent not installed"})
        if s.get("status") and str(s.get("status")).lower() not in ("provisioned", "running", "ok", "stopped"):
            health.append({"Entity": s.get("name"), "Type": "VM", "Severity": "Warning",
                           "Message": f"Status: {s.get('status')}"})

        hid = g(s, "parentServer", "id")
        hs = host_vm_stats.setdefault(hid, {"vms": 0, "vcpu": 0, "vram": 0, "on": 0})
        hs["vms"] += 1
        hs["vcpu"] += cores or 0
        hs["vram"] += to_int(mem) or 0
        hs["on"] += 1 if ps == "poweredOn" else 0

    # ---- Hosts
    for h in hosts:
        st = host_vm_stats.get(h.get("id"), {"vms": 0, "vcpu": 0, "vram": 0, "on": 0})
        cores = to_int(first(h.get("maxCores"), g(h, "hostInfo", "cpuCores")))
        mem = first(g(h, "stats", "maxMemory"), h.get("maxMemory"))
        used = first(g(h, "stats", "usedMemory"), h.get("usedMemory"))
        rows["vHost"].append({
            "Host": h.get("name"), "Cloud": cloud_name(h), "Cluster": cluster_of(h),
            "Status": h.get("status"), "Powerstate": power(h.get("powerState")),
            "IP Address": first(h.get("internalIp"), h.get("externalIp"), h.get("sshHost")),
            "CPU Model": strip(first(h.get("hardwareCpuModel"), g(h, "hostInfo", "cpuModel"), h.get("cpuModel"))),
            "CPU MHz": to_int(h.get("hardwareCpuFrequency")),
            "# Sockets": first(to_int(h.get("maxSockets")), to_int(h.get("maxCpu")),
                               to_int(g(h, "hostInfo", "cpuSockets"))),
            "# Cores": cores, "CPU Usage %": rnd(g(h, "stats", "cpuUsage")),
            "Memory MiB": mib(mem), "Memory Used MiB": mib(used), "Memory Usage %": pct(used, mem),
            "Storage MiB": mib(first(g(h, "stats", "maxStorage"), h.get("maxStorage"))),
            "Storage Used MiB": mib(first(g(h, "stats", "usedStorage"), h.get("usedStorage"))),
            "# NICs": len(h.get("interfaces") or []), "# VMs": st["vms"], "# VMs On": st["on"],
            "# vCPUs": st["vcpu"], "vCPUs per Core": round(st["vcpu"] / cores, 2) if cores else None,
            "vRAM MiB": mib(st["vram"]),
            "OS": first(g(h, "serverOs", "name"), h.get("platform"), h.get("osType")),
            "OS Version": h.get("platformVersion"),
            "Hypervisor": g(h, "computeServerType", "name"), "Agent Version": h.get("agentVersion"),
            "Vendor": strip(first(h.get("hardwareProductVendor"), h.get("manufacturer"),
                                  g(h, "hostInfo", "manufacturer"))),
            "Model": strip(first(h.get("hardwareProductName"), h.get("model"), g(h, "hostInfo", "model"))),
            "Maintenance": h.get("maintenanceMode"),
            "Creation date": parse_date(h.get("dateCreated")),
            "Host ID": h.get("id"), "External ID": h.get("externalId"), "UUID": h.get("uuid"),
            "Manager": manager})
        if h.get("status") and str(h.get("status")).lower() not in ("provisioned", "ok", "running"):
            health.append({"Entity": h.get("name"), "Type": "Host", "Severity": "Critical",
                           "Message": f"Host status: {h.get('status')}"})
        if h.get("maintenanceMode"):
            health.append({"Entity": h.get("name"), "Type": "Host", "Severity": "Info",
                           "Message": "Host is in maintenance mode"})
        if pct(used, mem) and pct(used, mem) > 90:
            health.append({"Entity": h.get("name"), "Type": "Host", "Severity": "Warning",
                           "Message": f"Memory usage {pct(used, mem)}%"})

    # ---- Clusters
    for cl in clusters:
        members = [m.get("id") if isinstance(m, dict) else m for m in cl.get("servers") or []]
        hs = [srv_by_id.get(m, {}) for m in members]
        vm_stats = [host_vm_stats.get(m, {}) for m in members]
        rows["vCluster"].append({
            "Cluster": cl.get("name"), "Type": g(cl, "type", "name"), "Layout": g(cl, "layout", "name"),
            "Cloud": first(g(cl, "zone", "name"), g(cl, "cloud", "name")),
            "Group": first(g(cl, "site", "name"), g(cl, "group", "name")),
            "Status": cl.get("status"), "Enabled": cl.get("enabled"), "# Hosts": len(members),
            "# VMs": sum(v.get("vms", 0) for v in vm_stats),
            "Total Cores": sum(to_int(h.get("maxCores")) or 0 for h in hs) or None,
            "Total Memory MiB": mib(sum(to_int(h.get("maxMemory")) or 0 for h in hs)) or None,
            "# vCPUs": sum(v.get("vcpu", 0) for v in vm_stats),
            "vRAM MiB": mib(sum(v.get("vram", 0) for v in vm_stats)),
            "Created": parse_date(cl.get("dateCreated")), "Cluster ID": cl.get("id")})

    # ---- Datastores
    ds_usage = {}
    for s in vms:
        for v in s.get("volumes") or []:
            dsid = first(v.get("datastoreId"), g(v, "datastore", "id"))
            u = ds_usage.setdefault(dsid, {"disks": 0, "vms": set(), "prov": 0})
            u["disks"] += 1
            u["vms"].add(s.get("id"))
            u["prov"] += to_int(v.get("maxStorage")) or 0
    cluster_hosts = {}
    for cl in clusters:
        cluster_hosts[cl.get("name")] = [
            name_of(srv_by_id.get(m.get("id") if isinstance(m, dict) else m))
            or (m.get("name") if isinstance(m, dict) else m)
            for m in cl.get("servers") or []]
    for d in inv.get("datastores", []):
        cfg = d.get("config") if isinstance(d.get("config"), dict) else {}
        cap = to_int(first(d.get("storageSize"), d.get("maxStorage"), d.get("capacity")))
        free = to_int(first(d.get("freeSpace"), d.get("free")))
        u = ds_usage.get(d.get("id"), {"disks": 0, "vms": set(), "prov": 0})
        free_pct = round(free / cap * 100, 1) if cap and free is not None else None
        d_clusters = d.get("_clusters") or ([name_of(d.get("cluster"))] if d.get("cluster") else [])
        d_hosts = sorted({h for cn in d_clusters for h in cluster_hosts.get(cn, []) if h})
        in_use = cap - free if cap and free is not None else None
        rows["vDatastore"].append({
            "Name": d.get("name"), "Display Name": d.get("displayName"),
            "Type": first(g(d, "datastoreType", "name"),
                          name_of(d.get("type")) if isinstance(d.get("type"), dict) else
                          (d.get("type") if d.get("type") != "generic" else None),
                          g(d, "datastoreType", "code"), d.get("externalType")),
            "Status": first(d.get("status"), "online" if d.get("online") else None),
            "Accessible": d.get("online"), "Active": d.get("active"),
            "Capacity MiB": mib(cap), "Provisioned MiB": mib(u["prov"]),
            "In Use MiB": mib(in_use), "Free MiB": mib(free), "Free %": free_pct,
            "% Provisioned": round(u["prov"] / cap * 100, 1) if cap else None,
            "# Disks": u["disks"], "# VMs": len(u["vms"]),
            "# Hosts": len(d_hosts) or None, "Hosts": ", ".join(d_hosts) or None,
            "Cluster": ", ".join(c for c in d_clusters if c) or None,
            "Cloud": first(g(d, "zone", "name"), g(d, "cloud", "name")),
            "Address": first(cfg.get("sourceHostname"), cfg.get("sourceHost"), cfg.get("host"),
                             cfg.get("serverHost"), cfg.get("target"), cfg.get("portal")),
            "Source": first(cfg.get("sourceDirPath"), cfg.get("exportFolder"), cfg.get("sourceDir"),
                            cfg.get("blockDevice")),
            "Path / URL": first(d.get("externalPath"), cfg.get("targetPath"), cfg.get("path"),
                                cfg.get("mountPath")),
            "Category": d.get("category"), "Visibility": d.get("visibility"),
            "Default": d.get("defaultStore"), "Allow Provision": d.get("allowProvision"),
            "Heartbeat Target": d.get("heartbeatTarget"),
            "Datastore ID": d.get("id"), "External ID": d.get("externalId")})
        if free_pct is not None and free_pct < ds_free_pct:
            health.append({"Entity": d.get("name"), "Type": "Datastore", "Severity": "Warning",
                           "Message": f"Only {free_pct}% free"})

    # ---- Networks
    nic_count = {}
    for r in rows["vNetwork"]:
        nic_count[r["Network"]] = nic_count.get(r["Network"], 0) + 1
    for n in inv.get("networks", []):
        rows["vNetworks"].append({
            "Network": n.get("name"), "Display Name": n.get("displayName"), "Type": g(n, "type", "name"),
            "Cloud": g(n, "zone", "name"), "CIDR": n.get("cidr"), "Gateway": n.get("gateway"),
            "DNS Primary": n.get("dnsPrimary"), "DNS Secondary": n.get("dnsSecondary"),
            "VLAN ID": n.get("vlanId"), "DHCP": n.get("dhcpServer"), "IP Pool": g(n, "pool", "name"),
            "Active": n.get("active"), "Visibility": n.get("visibility"),
            "# NICs": nic_count.get(n.get("name"), 0), "Network ID": n.get("id"),
            "External ID": n.get("externalId")})

    # ---- Instances
    for i in instances:
        srv_ids = [x.get("id") if isinstance(x, dict) else x for x in i.get("servers") or []]
        conn = ", ".join(str(ci.get("ip")) for ci in i.get("connectionInfo") or []
                         if ci.get("ip") not in (None, "", "0.0.0.0")) or None
        rows["vInstance"].append({
            "Instance": i.get("name"), "Display Name": i.get("displayName"), "Status": i.get("status"),
            "Instance Type": g(i, "instanceType", "name"), "Layout": g(i, "layout", "name"),
            "Plan": g(i, "plan", "name"), "Group": g(i, "group", "name"), "Cloud": g(i, "cloud", "name"),
            "Environment": i.get("environment"),
            "Created By": first(g(i, "createdBy", "username"), g(i, "createdBy", "name")),
            "Owner": g(i, "owner", "username"), "Created": parse_date(i.get("dateCreated")),
            "CPUs": to_int(i.get("maxCores")), "Memory MiB": mib(i.get("maxMemory")),
            "Storage MiB": mib(i.get("maxStorage")), "IP Addresses": conn, "# VMs": len(srv_ids),
            "VMs": ", ".join(str(name_of(srv_by_id.get(x)) or x) for x in srv_ids) or None,
            "Hostname": i.get("hostName"), "Tags": fmt_tags(i.get("tags")),
            "Labels": ", ".join(i.get("labels") or []) or None, "Description": i.get("description"),
            "Instance ID": i.get("id")})

    # ---- Snapshots
    inst_by_id = {str(i.get("id")): i for i in instances}
    # Age is measured at collection time, so a saved inventory (--from-json) reports
    # the same ages whenever it is processed.
    now = parse_date(inv.get("collected"))
    if not isinstance(now, dt.datetime):
        now = dt.datetime.now(dt.timezone.utc).replace(tzinfo=None)
    for iid, snaps in (inv.get("snapshots") or {}).items():
        inst = inst_by_id.get(iid, {})
        for sn in snaps:
            created = parse_date(first(sn.get("snapshotCreated"), sn.get("dateCreated")))
            age = (now - created).days if isinstance(created, dt.datetime) else None
            rows["vSnapshot"].append({
                "Instance": inst.get("name"), "Name": sn.get("name"), "Description": sn.get("description"),
                "Date / time": created, "Age (days)": age, "Type": sn.get("snapshotType"),
                "State": first(sn.get("state"), sn.get("status")),
                "Size MiB": mib(first(sn.get("maxStorage"),
                                      sum(to_int(f.get("fileSize")) or 0 for f in sn.get("snapshotFiles") or []))),
                "Memory": sn.get("memorySnapshot"), "Current": sn.get("currentlyActive"),
                "Cloud": g(inst, "cloud", "name"), "Snapshot ID": sn.get("id"),
                "External ID": sn.get("externalId"), "Instance ID": to_int(iid)})
            if age is not None and age > snapshot_age_days:
                health.append({"Entity": f"{inst.get('name')} / {sn.get('name')}", "Type": "Snapshot",
                               "Severity": "Warning", "Message": f"Snapshot is {age} days old"})

    # ---- Source
    ping, who = inv.get("ping") or {}, inv.get("whoami") or {}
    for k, v in [("Manager URL", manager), ("Build Version", ping.get("buildVersion")),
                 ("Appliance URL", ping.get("applianceUrl")),
                 ("Connected User", g(who, "user", "username")),
                 ("Tenant", g(who, "user", "account", "name")),
                 ("Collected (UTC)", parse_date(inv.get("collected"))),
                 ("Clouds", len(inv.get("clouds", []))), ("Clusters", len(clusters)),
                 ("Hosts", len(hosts)), ("VMs", len(vms)), ("Instances", len(instances)),
                 ("Datastores", len(inv.get("datastores", []))), ("Networks", len(inv.get("networks", []))),
                 ("Tool", f"HVMTools {__version__}")]:
        rows["vSource"].append({"Property": k, "Value": v})

    rows["vHealth"].sort(key=lambda r: ({"Critical": 0, "Warning": 1, "Info": 2}.get(r["Severity"], 3),
                                        {"Host": 0, "Datastore": 1, "Snapshot": 2, "VM": 3}.get(r["Type"], 4)))
    return rows


def flatten(obj, prefix="", out=None, depth=0):
    """Flatten nested dicts for raw sheets; lists become JSON strings."""
    out = {} if out is None else out
    for k, v in (obj or {}).items():
        key = f"{prefix}{k}"
        if isinstance(v, dict) and depth < 2:
            flatten(v, key + ".", out, depth + 1)
        elif isinstance(v, (list, dict)):
            out[key] = json.dumps(v, default=str)[:32000]
        else:
            out[key] = v
    return out


# --------------------------------------------------------------------------- Excel output
HEADER_FILL = PatternFill("solid", fgColor="1F4E78")
HEADER_FONT = Font(name="Arial", size=10, bold=True, color="FFFFFF")
BODY_FONT = Font(name="Arial", size=10)
SEV_FILL = {"Critical": PatternFill("solid", fgColor="F8CBAD"),
            "Warning": PatternFill("solid", fgColor="FFE699"),
            "Info": PatternFill("solid", fgColor="DDEBF7")}


def clean(v):
    if isinstance(v, str):
        return ILLEGAL_CHARACTERS_RE.sub("", v)
    if isinstance(v, (list, dict)):
        return json.dumps(v, default=str)
    return v


def write_sheet(wb, title, rows):
    ws = wb.create_sheet(title[:31])
    if not rows:
        ws["A1"] = "No data returned"
        ws["A1"].font = BODY_FONT
        return
    cols = []
    for r in rows:
        for k in r:
            if k not in cols:
                cols.append(k)
    ws.append(cols)
    for cell in ws[1]:
        cell.font, cell.fill = HEADER_FONT, HEADER_FILL
        cell.alignment = Alignment(vertical="center")
    widths = [len(c) + 3 for c in cols]  # room for the AutoFilter button
    for r in rows:
        ws.append([clean(r.get(c)) for c in cols])
        for i, c in enumerate(cols):
            v = r.get(c)
            widths[i] = max(widths[i], 16 if isinstance(v, dt.datetime) else len(str(v)) if v is not None else 0)
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.font = BODY_FONT
            if isinstance(cell.value, dt.datetime):
                cell.number_format = "yyyy-mm-dd hh:mm"
            elif isinstance(cell.value, int) and not isinstance(cell.value, bool) \
                    and not cols[cell.column - 1].endswith("ID"):
                cell.number_format = "#,##0"
        if title == "vHealth":
            sev = row[cols.index("Severity")].value
            if sev in SEV_FILL:
                for cell in row:
                    cell.fill = SEV_FILL[sev]
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = min(max(w + 2, 8), 60)
    ws.freeze_panes = "B2"
    ws.auto_filter.ref = ws.dimensions


def write_workbook(rows, path, inv=None, raw=False):
    wb = Workbook()
    wb.remove(wb.active)
    for title, data in rows.items():
        write_sheet(wb, title, data)
    if raw and inv:
        write_sheet(wb, "raw_servers", [flatten(s) for s in inv.get("servers", [])])
        write_sheet(wb, "raw_instances", [flatten(i) for i in inv.get("instances", [])])
        write_sheet(wb, "raw_datastores", [flatten(d) for d in inv.get("datastores", [])])
    wb.save(path)


# --------------------------------------------------------------------------- main
def main():
    ap = argparse.ArgumentParser(description="RVTools-style Excel inventory for HPE Morpheus VM Essentials")
    ap.add_argument("--version", action="version", version=f"HVMTools {__version__}")
    ap.add_argument("--url", default=os.getenv("MORPHEUS_URL"), help="Manager URL, e.g. https://hvm-mgr")
    ap.add_argument("-u", "--username", default=os.getenv("MORPHEUS_USER"))
    ap.add_argument("-p", "--password", default=os.getenv("MORPHEUS_PASS"), help="Prompted if omitted")
    ap.add_argument("--token", default=os.getenv("MORPHEUS_TOKEN"), help="API access token instead of user/pass")
    ap.add_argument("--insecure", action="store_true", help="Skip TLS verification (self-signed certs)")
    ap.add_argument("-o", "--output", help="Output .xlsx (default HVMTools_<host>_<date>.xlsx)")
    ap.add_argument("--no-snapshots", action="store_true", help="Skip per-instance snapshot queries")
    ap.add_argument("--no-stats", action="store_true", help="Skip per-server usage stats queries (used when the server list omits stats)")
    ap.add_argument("--raw", action="store_true", help="Add raw_servers / raw_instances / raw_datastores sheets with every API field")
    ap.add_argument("--dump-json", help="Also save the raw collected inventory to this JSON file")
    ap.add_argument("--from-json", help="Build the workbook from a previously dumped JSON (no API calls)")
    ap.add_argument("--snapshot-age", type=int, default=7, help="vHealth: flag snapshots older than N days")
    ap.add_argument("--page-size", type=int, default=100)
    ap.add_argument("--workers", type=int, default=8)
    a = ap.parse_args()

    if a.from_json:
        with open(a.from_json, encoding="utf-8") as f:
            inv = redact(json.load(f))
    else:
        if not a.url:
            ap.error("--url is required (or set MORPHEUS_URL)")
        c = MorpheusClient(a.url, verify=not a.insecure, page_size=a.page_size)
        if a.token:
            c.set_token(a.token)
        else:
            user = a.username or input("Username: ")
            pw = a.password or getpass.getpass(f"Password for {user}: ")
            c.login(user, pw)
        inv = collect(c, snapshots=not a.no_snapshots, stats=not a.no_stats, workers=a.workers)
        if a.dump_json:
            with open(a.dump_json, "w", encoding="utf-8") as f:
                json.dump(inv, f, indent=2, default=str)
            log(f"Raw inventory saved to {a.dump_json}")

    host = (inv.get("manager") or "hvm").split("//")[-1].split("/")[0].split(":")[0]
    out = a.output or f"HVMTools_{host}_{dt.datetime.now():%Y%m%d_%H%M}.xlsx"
    rows = build_sheets(inv, snapshot_age_days=a.snapshot_age)
    write_workbook(rows, out, inv, raw=a.raw)
    log(f"Wrote {out}: " + ", ".join(f"{k}={len(v)}" for k, v in rows.items() if k != "vSource"))


def run():
    try:
        main()
    except KeyboardInterrupt:
        raise SystemExit("Interrupted.")
    except requests.exceptions.SSLError as e:
        raise SystemExit(f"TLS error: {e}\nIf the manager uses a self-signed certificate, add --insecure.")
    except requests.exceptions.ConnectionError as e:
        raise SystemExit(f"Cannot connect to the manager: {e}\nCheck --url, DNS and that port 443 is reachable.")
    except requests.exceptions.Timeout:
        raise SystemExit("The manager did not respond in time.")
    except requests.exceptions.HTTPError as e:
        raise SystemExit(f"API error: {e}")
    except PermissionError as e:
        raise SystemExit(f"Cannot access {e.filename}: permission denied (if it's the workbook, is it open in Excel?)")
    except (OSError, json.JSONDecodeError) as e:
        raise SystemExit(f"Error: {e}")


if __name__ == "__main__":
    run()
